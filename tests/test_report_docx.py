"""Testy modulu report_docx (etapa M1) – investiční report DOCX.

Reference (tests/fixtures/report_docx_reference.json) jsou otisky všech částí souboru DOCX pro tři
scénáře; při jejím vytvoření byla ověřena bajtová shoda s DOCX z původního kódu (main 5fef75c).
Po ZÁMĚRNÉ a schválené změně reportu ji obnovte: python tests/test_report_docx.py --update
"""

import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
import investice as inv  # noqa: E402
import report_docx as rd  # noqa: E402
from test_investice import CASES, machines  # noqa: E402

REFERENCE = ROOT / "tests" / "fixtures" / "report_docx_reference.json"
REPORT_CASES = ("pesimisticky", "optimisticky", "oba_ztratove")


def inputs(case: str) -> dict:
    """Vstupy reportu stejně, jak je předává app.py (hodnoty z investice.evaluate a postranního panelu)."""
    m1, m2, scenario, fixed, days, n = CASES[case]
    df = machines()
    row_1 = df[df["machine_type"] == m1].iloc[0]
    row_2 = df[df["machine_type"] == m2].iloc[0]
    lo, hi = inv.scenario_days(scenario)
    r = inv.evaluate(row_1, row_2, m1, m2, fixed, lo, hi, days, n)
    keys = ("better_machine", "break_even_dumper", "break_even_mini", "expected_profit_dumper", "expected_profit_mini",
            "p10_dumper", "p10_mini", "p50_dumper", "p50_mini", "p90_dumper", "p90_mini", "required_price_1",
            "required_price_2", "risk_loss_dumper", "risk_loss_mini", "scorecard_dumper", "scorecard_mini")
    return dict({k: r[k] for k in keys}, machine_1=m1, machine_2=m2, fixed_cost=fixed, scenario=scenario,
                min_days=lo, max_days=hi, n=n, expected_days=days,
                avg_daily_mini=float(row_1["daily_price"]), avg_daily_dumper=float(row_2["daily_price"]))


def parts(data: bytes) -> dict:
    z = zipfile.ZipFile(io.BytesIO(data))
    return {name: hashlib.sha256(z.read(name)).hexdigest() for name in sorted(z.namelist())}


def fingerprint(data: bytes) -> list:
    """Otisk DOCX bez časových značek ZIP: pořadí položek, jejich názvy, komprese a obsah.

    Celé bajty souboru porovnávat nelze – zipfile dá každé položce aktuální čas (přesnost 2 s),
    takže dvě jinak shodná sestavení se v rychlém sledu občas liší jen časovou značkou.
    """
    z = zipfile.ZipFile(io.BytesIO(data))
    return [(i.filename, i.compress_type, hashlib.sha256(z.read(i)).hexdigest()) for i in z.infolist()]


def rezip(data: bytes, date_time=None, replace=None) -> bytes:
    """Kopie DOCX s jinou časovou značkou položek nebo s nahrazeným obsahem jedné části (jen pro testy)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            blob = src.read(info)
            if replace and info.filename == replace[0]:
                blob = replace[1](blob)
            new = zipfile.ZipInfo(info.filename, date_time=date_time or info.date_time)
            new.compress_type = info.compress_type
            dst.writestr(new, blob)
    return out.getvalue()


def text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])


@pytest.mark.parametrize("case", REPORT_CASES)
def test_report_matches_reference_exactly(case):
    ref = json.loads(REFERENCE.read_text(encoding="utf-8"))[case]
    assert parts(rd.build_report(**inputs(case))) == ref


def test_report_is_deterministic():
    kw = inputs("pesimisticky")
    a, b = rd.build_report(**kw), rd.build_report(**kw)
    assert fingerprint(a) == fingerprint(b)
    assert text(a) == text(b)


def test_fingerprint_ignores_only_zip_timestamps():
    data = rd.build_report(**inputs("pesimisticky"))
    shifted = rezip(data, date_time=(2001, 2, 3, 4, 5, 6))
    assert shifted != data                                          # bajty se liší jen časovou značkou
    assert fingerprint(shifted) == fingerprint(data)


def test_fingerprint_detects_content_changes():
    kw = inputs("pesimisticky")
    data = rd.build_report(**kw)
    changed = rezip(data, replace=("word/document.xml", lambda b: b.replace("Mini bagr".encode(), "Mini bágr".encode(), 1)))
    assert fingerprint(changed) != fingerprint(data)                # změna textu v dokumentu
    for case in ("optimisticky", "oba_ztratove"):                   # jiné vstupy → jiný report
        assert fingerprint(rd.build_report(**inputs(case))) != fingerprint(data)
    kw["fixed_cost"] = kw["fixed_cost"] + 1
    assert fingerprint(rd.build_report(**kw)) != fingerprint(data)  # i drobná změna jedné hodnoty


def test_report_content_and_branches():
    t = text(rd.build_report(**inputs("vychozi")))
    assert "Decision Engine – investiční report" in t and "Mini bagr" in t and "Pásový dumper" in t
    assert "Doporučený stroj: Pásový dumper." in t and "Scénář trhu: Realistický (15–25 dní pronájmu měsíčně)" in t
    for case in ("pesimisticky", "oba_ztratove"):                   # v obou scénářích prodělávají oba stroje
        assert "Žádný stroj není ziskový." in text(rd.build_report(**inputs(case)))
    kw = inputs("vychozi")                                          # shoda kladných zisků (ztráta má přednost)
    kw["expected_profit_dumper"] = kw["expected_profit_mini"]
    assert "Stroje vycházejí stejně." in text(rd.build_report(**kw))


def test_report_requires_all_inputs():
    kw = inputs("pesimisticky")
    kw.pop("scenario")
    with pytest.raises(TypeError):
        rd.build_report(**kw)


def test_module_has_no_streamlit_dependency():
    import ast
    tree = ast.parse((ROOT / "report_docx.py").read_text(encoding="utf-8"))
    imported = {n.module.split(".")[0] for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) and n.module}
    imported |= {a.name.split(".")[0] for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names}
    assert "streamlit" not in imported, imported


if __name__ == "__main__":
    if "--update" in sys.argv:
        REFERENCE.write_text(json.dumps({c: parts(rd.build_report(**inputs(c))) for c in REPORT_CASES},
                                        ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")
        print(f"Reference uložena: {REFERENCE}")
    else:
        print(__doc__)
