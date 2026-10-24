"""Investiční report ve formátu DOCX (etapa M1).

Funkce je přesunutá z app.py (dříve build_docx_report) BEZ ZMĚNY obsahu dokumentu: stejné texty,
tabulky, styly i pořadí. Vstupy jsou hodnoty, které dříve četla z proměnných app.py; app.py je
předává pojmenovaně (stejná jména). Modul nezávisí na Streamlitu.
"""

from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor

from formatting import fmt_czk, fmt_days, fmt_num, fmt_pct


def build_report(*, avg_daily_dumper, avg_daily_mini, better_machine, break_even_dumper, break_even_mini,
                 expected_days, expected_profit_dumper, expected_profit_mini, fixed_cost, machine_1,
                 machine_2, max_days, min_days, n, p10_dumper, p10_mini, p50_dumper, p50_mini, p90_dumper,
                 p90_mini, required_price_1, required_price_2, risk_loss_dumper, risk_loss_mini, scenario,
                 scorecard_dumper, scorecard_mini) -> bytes:
    """Vytvoří DOCX a vrátí ho jako bajty (stejný obsah jako tlačítko „Stáhnout report“)."""
    doc = Document()

    # Základní písmo dokumentu
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        heading = doc.styles[style_name]
        heading.font.name = "Aptos Display"
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(20, 48, 79)

    title = doc.add_heading("Decision Engine – investiční report", level=1)
    title_run = title.runs[0]
    title_run.font.size = Pt(20)

    doc.add_paragraph(f"Porovnání strojů: {machine_1} a {machine_2}")
    demo = doc.add_paragraph()
    demo_run = demo.add_run(
        "Ukázkový výstup pilotního prototypu. Všechny vstupy i výsledky vycházejí "
        "z fiktivních demonstračních dat."
    )
    demo_run.italic = True
    demo_run.font.color.rgb = RGBColor(122, 74, 6)

    doc.add_heading("Zadané předpoklady", level=1)
    doc.add_paragraph(f"Fixní náklad na stroj: {fmt_czk(fixed_cost)} měsíčně")
    doc.add_paragraph(f"Scénář trhu: {scenario} ({min_days}–{max_days} dní pronájmu měsíčně)")
    doc.add_paragraph(f"Počet simulovaných měsíců: {fmt_num(n)}")
    doc.add_paragraph(f"Odhad poptávky pro rychlý výpočet: {expected_days} dní/měsíc")

    doc.add_heading("Klíčové investiční ukazatele", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Ukazatel"
    hdr[1].text = machine_1
    hdr[2].text = machine_2

    rows = [
        ("Denní cena pronájmu", fmt_czk(avg_daily_mini), fmt_czk(avg_daily_dumper)),
        ("Bod zvratu (jen fixní náklady)", fmt_days(break_even_mini, 2), fmt_days(break_even_dumper, 2)),
        ("Šance ztrátového měsíce", fmt_pct(risk_loss_mini), fmt_pct(risk_loss_dumper)),
        ("Očekávaný měsíční zisk (simulace)", fmt_czk(expected_profit_mini), fmt_czk(expected_profit_dumper)),
    ]
    for metric, v1, v2 in rows:
        cells = table.add_row().cells
        cells[0].text = metric
        cells[1].text = v1
        cells[2].text = v2

    # Stejné tři případy jako hlavní karta v záložce Doporučení.
    doc.add_heading("Doporučení", level=1)
    verdict_p = doc.add_paragraph()
    if expected_profit_mini < 0 and expected_profit_dumper < 0:
        verdict_p.add_run("Žádný stroj není ziskový. ").bold = True
        verdict_p.add_run(
            f"Relativně lépe vychází {better_machine}, ale při zadaném nastavení jsou obě "
            "varianty v průměru ztrátové. Investici nedoporučujeme realizovat bez zvýšení ceny "
            "pronájmu, vyššího vytížení nebo snížení nákladů."
        )
    elif expected_profit_mini == expected_profit_dumper:
        verdict_p.add_run("Stroje vycházejí stejně. ").bold = True
        verdict_p.add_run("Oba mají stejný očekávaný zisk – rozhodněte podle rizika a ceny.")
    else:
        verdict_p.add_run(f"Doporučený stroj: {better_machine}. ").bold = True
        verdict_p.add_run(
            "Má vyšší očekávaný měsíční zisk v simulaci. Podrobné srovnání rizika a bodu "
            "zvratu je v tabulce výše."
        )
    doc.add_paragraph(
        f"Doplňkové bodové hodnocení (zisk, šance ztráty, velmi nepříznivý měsíc, návratnost; "
        f"při shodě bez bodu): {machine_1} {scorecard_mini} / 4, {machine_2} {scorecard_dumper} / 4."
    )

    doc.add_heading("Rozpětí možných výsledků (simulace)", level=1)
    doc.add_paragraph(
        f"Simulace vytváří {fmt_num(n)} možných měsíců s různou poptávkou, servisními náklady, "
        "poruchami a cenou. Následující hodnoty shrnují rozpětí měsíčního výsledku."
    )
    doc.add_paragraph(
        f"Slabý měsíc (P10 – jen 1 z 10 měsíců je horší): {machine_1} {fmt_czk(p10_mini)}, "
        f"{machine_2} {fmt_czk(p10_dumper)}"
    )
    doc.add_paragraph(
        f"Typický měsíc (medián): {machine_1} {fmt_czk(p50_mini)}, {machine_2} {fmt_czk(p50_dumper)}"
    )
    doc.add_paragraph(
        f"Dobrý měsíc (P90 – jen 1 z 10 měsíců je lepší): {machine_1} {fmt_czk(p90_mini)}, "
        f"{machine_2} {fmt_czk(p90_dumper)}"
    )

    price_heading = doc.add_heading("Doporučená minimální cena pronájmu", level=1)
    price_heading.paragraph_format.page_break_before = True

    price_table = doc.add_table(rows=1, cols=3)
    price_table.style = "Table Grid"
    header_cells = price_table.rows[0].cells
    header_cells[0].text = "Stroj"
    header_cells[1].text = "Minimální cena"
    header_cells[2].text = "Hodnocení aktuální ceny"

    for name, required, actual in (
        (machine_1, required_price_1, avg_daily_mini),
        (machine_2, required_price_2, avg_daily_dumper),
    ):
        row = price_table.add_row().cells
        row[0].text = str(name)
        row[1].text = f"{fmt_czk(required)}/den"
        row[2].text = "Cena je dostačující" if actual >= required else "Cena je pod minimem pro zisk"

    doc.add_paragraph(
        f"Minimální cena ukazuje, při jaké denní sazbě se při poptávce {expected_days} dní "
        "měsíčně pokryjí fixní a servisní náklady."
    )

    doc.add_heading("Interpretace výsledků", level=1)
    doc.add_paragraph(
        "Model simuluje možné scénáře poptávky po strojích a zohledňuje náklady na servis, "
        "fixní náklady a pravděpodobnost poruch. Výsledkem je odhad očekávaného měsíčního "
        "zisku a rizika ztráty."
    )
    doc.add_paragraph(
        "Výběr nákupní sestavy (záložka Nákup strojů) a plán zakázek (záložka Plán zakázek) "
        "jsou v aplikaci samostatné analýzy na ukázkových datech. Tento report je nezahrnuje "
        "a navzájem zatím propojené nejsou."
    )
    doc.add_paragraph(
        "Doporučení jsou sestavena pevně danými pravidly z vypočtených ukazatelů."
    )

    doc.add_heading("Doporučený další postup", level=1)
    for step in [
        f"Ověřit skutečnou měsíční poptávku po stroji {better_machine}.",
        "Porovnat aktuální cenu pronájmu s vypočtenou minimální denní sazbou.",
        "Doplnit reálné servisní náklady, odstávky, pojištění a sezónnost.",
        "Znovu vyhodnotit investici v realistickém scénáři.",
    ]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Podmínky ziskovosti", level=1)
    doc.add_paragraph(
        "Investice by byla výrazně atraktivnější při splnění některých z následujících podmínek:"
    )
    for item in [
        "zvýšení denní ceny pronájmu",
        "vyšší vytížení strojů (více dní pronájmu za měsíc)",
        "snížení fixních nebo servisních nákladů",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()
